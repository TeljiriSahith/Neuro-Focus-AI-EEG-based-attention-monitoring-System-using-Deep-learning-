from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    
    # Define the slide content (Shortened for 2-3 minute pacing)
    slides_data = [
        {
            "title": "NEUROFOCUS: EEG-BASED ATTENTION MONITORING",
            "content": [
                "Deep Learning & Single-Channel EEG",
                "Department: CSE (Data Science)",
                "Institution: Geethanjali College of Engineering and Technology",
                "Presented By: Sahith Yadav & Team"
            ]
        },
        {
            "title": "Introduction",
            "content": [
                "The Goal: Monitor attention to improve productivity and learning.",
                "The Problem: Traditional surveys are subjective and not real-time.",
                "The EEG Solution: Brainwaves provide continuous cognitive data.",
                "The Challenge: Clinical multi-channel EEG is bulky and expensive.",
                "Our Approach: NeuroFocus AI uses a single-channel EEG and a 1D-CNN for automated monitoring."
            ]
        },
        {
            "title": "Domain & Technologies",
            "content": [
                "Domains: Artificial Intelligence, Deep Learning, Brain-Computer Interfaces (BCI).",
                "Core Language: Python",
                "Deep Learning Framework: PyTorch",
                "Data & Signal Processing: Pandas, SciPy, NumPy",
                "UI & Visualization: Streamlit, Matplotlib"
            ]
        },
        {
            "title": "Problem Statement",
            "content": [
                "Subjective Methods: Self-reporting lacks objective, continuous data.",
                "Noisy Signals: Raw EEG is easily distorted by blinks and movement.",
                "Complex Extraction: Traditional models require difficult manual feature engineering.",
                "Heavy Hardware: Multi-channel caps are impractical for daily use.",
                "No Real-Time Systems: Most current models are strictly offline."
            ]
        },
        {
            "title": "About NeuroFocus AI",
            "content": [
                "Single-Channel Input: Captures signals strictly from the FP1 (frontal) electrode.",
                "Deep Learning: Uses a 1D-CNN to skip manual feature extraction.",
                "Optimized Threshold: Uses a 0.35 classification threshold for high accuracy.",
                "Smart Aggregation: Averages 1-second predictions into stable 10-second trends.",
                "Interactive UI: Provides live dashboards and exportable PDF reports."
            ]
        },
        {
            "title": "Objectives & SDG Mapping",
            "content": [
                "Key Objectives: Preprocess single-channel EEG, build a 1D-CNN, and deploy a web UI.",
                "SDG 3 (Health): Early identification of cognitive overload and fatigue.",
                "SDG 4 (Education): Enhances adaptive e-learning via engagement tracking.",
                "SDG 9 (Innovation): Promotes accessible Brain-Computer Interfaces."
            ]
        },
        {
            "title": "System Architecture",
            "content": [
                "1. Input: Single FP1 Channel or CSV Upload.",
                "2. Preprocessing: 250Hz Resample, Z-Score Normalization, 1s Windows.",
                "3. Model: 1D-CNN (Conv1D + MaxPool + Dropout).",
                "4. Output: Sigmoid Probability & 10s Aggregation.",
                "5. UI: Streamlit Dashboard & Session Reports."
            ]
        },
        {
            "title": "Literature Survey",
            "content": [
                "CNN on EEG (Cecotti, 2011): Deep learning extracts features well, but limited to offline tasks.",
                "EEGNet (Lawhern, 2018): Compact networks generalize well, but need multi-channel setups.",
                "CNN/RNN (Roy, 2019): Outperforms handcrafted features, but requires high compute.",
                "NeuroFocus (Ours): Achieves real-time, single-channel tracking for practical use."
            ]
        },
        {
            "title": "Design — Preprocessing",
            "content": [
                "Resampling: Converts raw 512Hz signals to a standardized 250Hz.",
                "Segmentation: Chops data into exact 1-second chunks (250 samples).",
                "Normalization: Applies Z-Score scaling to handle amplitude variations.",
                "Tensor Conversion: Formats data into PyTorch tensors for the 1D-CNN."
            ]
        },
        {
            "title": "Design — 1D-CNN Model",
            "content": [
                "Architecture: 3 Convolutional layers followed by Max Pooling.",
                "Activation: ReLU activation and Batch Normalization.",
                "Classifier: Flatten -> Dropout (0.5) -> Linear Layer.",
                "Training Details: BCEWithLogitsLoss, Adam Optimizer, 30 Epochs."
            ]
        },
        {
            "title": "UI Execution (Upload & Processing)",
            "content": [
                "File Verification: Ensures the CSV contains the required FP1 column.",
                "Live Metadata: Displays total samples, duration in seconds, and total windows.",
                "Automated Pipeline: Normalizes, creates windows, and runs the CNN instantly."
            ]
        },
        {
            "title": "UI Execution (Summary Dashboard)",
            "content": [
                "KPI Dashboard: Shows Average, Peak, and Lowest Attention percentages.",
                "Session Stats: Displays exact attentive vs. inattentive seconds.",
                "AI Interpretation: Provides an automated textual summary of the user's focus."
            ]
        },
        {
            "title": "UI Execution (Visual Analytics)",
            "content": [
                "Trend Line: Tracks 10-second moving averages over time.",
                "Scatter Plot: Highlights exact focus percentage at specific intervals.",
                "Pie Chart: Compares overall attentive vs. inattentive duration."
            ]
        },
        {
            "title": "Model Performance & Thresholds",
            "content": [
                "Why 0.35 Threshold? Selected to balance precision and recall perfectly.",
                "Threshold 0.35: 84.04% Accuracy | 85.53% F1-Score (Selected)",
                "Threshold 0.40: 83.68% Accuracy | 84.90% F1-Score",
                "Threshold 0.50: 82.05% Accuracy | 82.51% F1-Score"
            ]
        },
        {
            "title": "Verification & Testing",
            "content": [
                "Confusion Matrix Results:",
                "True Positives (Attentive): 520",
                "True Negatives (Inattentive): 407",
                "False Positives: 83 | False Negatives: 93",
                "Testing Phase: All unit tests (CSV loading, model prediction) and functional UI tests passed."
            ]
        },
        {
            "title": "Conclusion & Future Scope",
            "content": [
                "Success: Single-channel FP1 EEG is highly effective for attention tracking.",
                "High Accuracy: Reached 84.04% Accuracy and 85.53% F1-Score.",
                "Practical Tool: The web dashboard is ready for educational and workplace deployment.",
                "Future Scope: Adding live Bluetooth hardware integration and audio-visual alerts."
            ]
        },
        {
            "title": "Thank You",
            "content": [
                "Project: NeuroFocus AI",
                "Questions & Discussion"
            ]
        }
    ]

    # Title Slide Layout (Layout 0)
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = slides_data[0]["title"]
    subtitle.text = "\n".join(slides_data[0]["content"])

    # Bullet Slide Layout (Layout 1)
    bullet_slide_layout = prs.slide_layouts[1]

    # Generate remaining slides
    for slide_data in slides_data[1:]:
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        title_shape.text = slide_data["title"]
        
        tf = body_shape.text_frame
        tf.text = slide_data["content"][0]
        
        for point in slide_data["content"][1:]:
            p = tf.add_paragraph()
            p.text = point
            p.level = 0

    # Save the presentation
    prs.save('NeuroFocus_AI_Review.pptx')
    print("Presentation generated successfully: NeuroFocus_AI_Review.pptx")

if __name__ == "__main__":
    create_presentation()